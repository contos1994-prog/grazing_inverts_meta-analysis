"""
Create MetaAnalysis_data_v2.xlsx from the original workbook.

- Fixes drag-fill (auto-increment) errors
- Standardises exclusion / sampling dates into clean numeric columns
- Adds a small constant to zero means (and their zero SEs)
- Adds cleaned moderator columns used by the R analysis
- Highlights every edited cell in yellow, derived columns in blue,
  and records each change in a 'Cleaning_log' sheet

The original file is left untouched.
Run from the repo root:  python3 scripts/00_make_v2.py
"""
import re
import datetime as dt
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

SRC = "MetaAnalysis_data.xlsx"
DST = "MetaAnalysis_data_v2.xlsx"

EDIT = PatternFill("solid", fgColor="FFF2A8")     # yellow = value changed
DERIVED = PatternFill("solid", fgColor="D9E8FB")  # blue   = new derived column

wb = load_workbook(SRC)
log = []


def header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def col_by_prefix(hmap, prefix):
    hits = [c for h, c in hmap.items() if h.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return hits[0]


def set_cell(ws, row, col, new, reason, sheet):
    cell = ws.cell(row=row, column=col)
    old = cell.value
    if old == new:
        return
    cell.value = new
    cell.fill = EDIT
    log.append([sheet, row, ws.cell(row=1, column=col).value, repr(old), repr(new), reason])


# ---------------------------------------------------------------- helpers
YEAR_RE = re.compile(r"(?<!\d)(19\d{2}|20\d{2})(?!\d)")


def excel_serial_year_bug(v):
    """A year typed into a date-formatted cell is stored as serial N -> 1905-xx-xx."""
    if isinstance(v, dt.datetime) and v.year in (1904, 1905):
        return (v - dt.datetime(1899, 12, 30)).days
    return None


def to_year(v):
    if v is None:
        return None
    y = excel_serial_year_bug(v)
    if y is not None:
        return y
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year
    if isinstance(v, (int, float)):
        return int(v) if 1800 < v < 2100 else None
    m = YEAR_RE.findall(str(v))
    return int(m[0]) if m else None


def duration_years(text):
    """Parse explicit durations like '24 years', '3-5 years', '>40 years'.
    Ranges -> midpoint; open-ended bounds ('>40', 'at least 10') -> the bound."""
    if text is None or isinstance(text, (dt.datetime, int, float)):
        return None, None
    s = str(text).lower()
    if "year" not in s:
        return None, None
    m = re.search(r"(\d+(?:\.\d+)?)\s*[-–to]+\s*(\d+(?:\.\d+)?)\s*year", s)
    if m:
        a, b = float(m.group(1)), float(m.group(2))
        return (a + b) / 2, f"range {a:g}-{b:g} yr -> midpoint"
    m = re.search(r"(\d+(?:\.\d+)?)\s*year", s)
    if m:
        v = float(m.group(1))
        note = "open-ended bound used" if (">" in s or "at least" in s) else "stated duration"
        return v, note
    return None, None


def sampling_years(v):
    if v is None:
        return None, None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year, v.year
    if isinstance(v, (int, float)):
        return int(v), int(v)
    ys = [int(y) for y in YEAR_RE.findall(str(v))]
    if not ys:
        return None, None
    # "2003-2011" style ranges are covered because both ends are captured
    return min(ys), max(ys)


# ================================================================ INVERTEBRATES
ws = wb["Invertebrates"]
S = "Invertebrates"
H = header_map(ws)
c_id = H["Study_ID"]
c_auth = H["Author_Year"]
c_dur = H["Grazing_Exclusion_Duration"]
c_dury = H["Grazing_Exclusion_Duration_Year"]
c_samp = H["Sampling_date"]
c_herb = col_by_prefix(H, "Herbivore_Type")
c_size = col_by_prefix(H, "Grazer_Size_Class_Standardized")
c_meas = col_by_prefix(H, "Measured")
c_strat = col_by_prefix(H, "type of invertebrates")
c_mg, c_sg, c_ng = H["Mean_Grazed"], H["SD_or_SE_Grazed"], H["n_Grazed"]
c_mc, c_sc, c_nc = H["Mean_Control"], H["SD_or_SE_Control"], H["n_Control"]
last = ws.max_row
while ws.cell(row=last, column=c_id).value in (None, ""):
    last -= 1
rows = range(2, last + 1)

# --- 1. whitespace in Study_ID / Author_Year, text-stored numbers
for r in rows:
    for c in (c_id, c_auth):
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str) and v != v.strip().replace("\xa0", ""):
            set_cell(ws, r, c, v.replace("\xa0", " ").strip(), "strip whitespace", S)
    for c in (c_mg, c_sg, c_ng, c_mc, c_sc, c_nc):
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str):
            try:
                set_cell(ws, r, c, float(v), "number stored as text", S)
            except ValueError:
                pass


# --- 1b. standardise Order (spelling, synonyms, wrong rank) and Trophic_Level
c_order, c_fam = H["Order"], H["Family"]
c_gen, c_common = H["Genus or species"], H["Common name"]
c_troph, c_group = H["Trophic_Level"], col_by_prefix(H, "Group")

ORDER_FIX = {
    "Aranaea": ("Araneae", "misspelling"),
    "Spider": ("Araneae", "common name entered as order"),
    "Hymenpotera": ("Hymenoptera", "misspelling"),
    "Homeptera": ("Hemiptera", "misspelling of Homoptera; Homoptera is now part of Hemiptera"),
    "Homoptera": ("Hemiptera", "obsolete order; Homoptera is now part of Hemiptera"),
    "Isoptera (Rhinotermitidae, Termitidae)": ("Isoptera", "families removed (already in Family column)"),
}
# family (or suborder) entered as order -> (order, family to fill if Family is blank)
ORDER_RANK = {
    "Formicidae": ("Hymenoptera", "Formicidae"),
    "Lumbricidae": ("Crassiclitellata", "Lumbricidae"),
    "Enchytraeidae": ("Enchytraeida", "Enchytraeidae"),
    "Oribatida": ("Sarcoptiformes", "Oribatida"),   # as for #1021/#1022
}
# blank Order filled where the common name / group identifies it -> (order, family)
ORDER_FILL = {
    "spider": ("Araneae", None),
    "ground beetle": ("Coleoptera", "Carabidae"),
    "plant hopper": ("Hemiptera", None),
    "nematodes": ("Nematoda", None),
    "mites": ("Acari", None),
    "enchytraeids": ("Enchytraeida", "Enchytraeidae"),
    "molluscs (snails/slugs)": ("Gastropoda", None),
}

for r in rows:
    v = ws.cell(row=r, column=c_order).value
    s = v.replace("\xa0", " ").strip() if isinstance(v, str) else v
    fam = ws.cell(row=r, column=c_fam).value
    common = str(ws.cell(row=r, column=c_common).value or "").strip().lower()
    group = str(ws.cell(row=r, column=c_group).value or "").strip()
    new, new_fam, why = s, None, "Order: strip whitespace"
    if s in ORDER_FIX:
        new, why = ORDER_FIX[s][0], f"Order: {ORDER_FIX[s][1]}"
    elif s in ORDER_RANK:
        new, new_fam = ORDER_RANK[s]
        why = f"Order: '{s}' is a family/suborder, not an order"
    elif s == "Oligochaeta" and "Lumbricidae" in str(fam or ""):
        new, why = "Crassiclitellata", "Order: Oligochaeta is a subclass; Lumbricidae + Megascolecidae are Crassiclitellata"
    elif s == "Gastropoda" and "Helminthoglypta" in str(ws.cell(row=r, column=c_gen).value or ""):
        new, why = "Stylommatophora", "Order: Gastropoda is a class; Helminthoglypta is Stylommatophora (as in #2176)"
    elif s is None and common in ORDER_FILL:
        new, new_fam = ORDER_FILL[common]
        why = f"Order: blank, filled from common name '{common}'"
    elif s is None and group == "Nematoda":
        new, why = "Nematoda", "Order: blank, filled from Group 'Nematoda'"
    set_cell(ws, r, c_order, new, why, S)
    if new_fam and fam in (None, ""):
        set_cell(ws, r, c_fam, new_fam, why.replace("Order:", "Family filled -"), S)

TROPHIC = {
    "herbivore": "Herbivore", "herbivores": "Herbivore",
    "phytophagus": "Herbivore", "plant_feeders": "Herbivore",
    "predator": "Predator", "predators": "Predator", "predatory": "Predator",
    "carnivore": "Predator", "carnivores": "Predator", "zoophagus": "Predator",
    "detritivore": "Detritivore", "detritivores": "Detritivore", "decomposers": "Detritivore",
    "bacterial feeders": "Bacterivore", "bacterivores": "Bacterivore",
    "detritivores - bacterial feeders": "Bacterivore", "detritivores -bacterivores": "Bacterivore",
    "fungal feeders": "Fungivore", "fungivores": "Fungivore",
    "detritivores -fungal feeders": "Fungivore", "detritivores - fungivores": "Fungivore",
    "algal feeders": "Algivore",
    "omnivore": "Omnivore", "omnivores": "Omnivore",
    "parasitoid": "Parasitoid",
    "pollinator": "Pollinator", "pollinators": "Pollinator",
    "scavenger": "Scavenger",
    "detritivore/herbivore": "Herbivore/detritivore", "herbivore or detritivore": "Herbivore/detritivore",
    "total": "Total",
}
for r in rows:
    v = ws.cell(row=r, column=c_troph).value
    if v is None:
        if ws.cell(row=r, column=c_order).value == "Araneae":
            set_cell(ws, r, c_troph, "Predator", "Trophic_Level: blank, spiders are predators", S)
        continue
    key = " ".join(str(v).split()).lower()
    new = TROPHIC[key]   # KeyError = new label in the data, add it to TROPHIC
    set_cell(ws, r, c_troph, new, f"Trophic_Level: standardised '{v}' -> '{new}'", S)

# formula cells: openpyxl cannot keep their cached values, so replace with values.
# AH490 (#3610 control SE) is '=AG491-AG490', i.e. it subtracts this row's control
# mean from the NEXT study's control mean -> broken; blank it so the row drops.
wb_vals = load_workbook(SRC, data_only=True)[S]
for r in rows:
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str) and v.startswith("="):
            if (r, c) == (490, c_sc):
                set_cell(ws, r, c, None, f"broken formula {v} references another study's mean - SE needs re-extracting", S)
            else:
                set_cell(ws, r, c, wb_vals.cell(row=r, column=c).value, f"formula {v} replaced by its value", S)


def rows_of(study):
    return [r for r in rows if ws.cell(row=r, column=c_id).value == study]


# --- 2. drag-fill fixes: copy first row of the block down
def fill_down(study, cols, reason):
    rr = rows_of(study)
    for c in cols:
        first = ws.cell(row=rr[0], column=c).value
        for r in rr[1:]:
            set_cell(ws, r, c, first, reason, S)


fill_down("#1013", [c_dur, c_dury, c_samp], "drag-fill error: exclusion year and sampling date auto-incremented")
fill_down("#2260", [c_samp], "drag-fill error: sampling date auto-incremented")
fill_down("#1589", [c_dur, c_dury, c_samp], "drag-fill error: exclusion year and sampling date auto-incremented")
fill_down("#1289", [c_auth], "drag-fill error: publication year auto-incremented")
fill_down("#2492", [c_auth], "drag-fill error: publication year auto-incremented")

# --- 3. Excel serial-number year bug (e.g. 1999 shown as 1905-06-21)
for r in rows:
    v = ws.cell(row=r, column=c_dury).value
    y = excel_serial_year_bug(v)
    if y is not None:
        set_cell(ws, r, c_dury, y, "year stored as Excel date serial (displayed as 1905)", S)

# --- 4. zero means / zero SEs
# lnRR is undefined for a zero mean. Add a small constant c to BOTH group means
# of the affected row (keeps the ratio symmetric), and set a zero SE to c
# (a zero SE would give the effect near-infinite weight).
# c = half the smallest non-zero mean reported in that study.
zero_rows = [r for r in rows if (ws.cell(row=r, column=c_mg).value == 0
                                 or ws.cell(row=r, column=c_mc).value == 0)]
study_c = {}
for r in zero_rows:   # compute c from the original values before any edits
    study = ws.cell(row=r, column=c_id).value
    means = [ws.cell(row=q, column=c).value for q in rows_of(study) for c in (c_mg, c_mc)]
    study_c[study] = min(m for m in means if isinstance(m, (int, float)) and m > 0) / 2
for r in zero_rows:
    c = study_c[ws.cell(row=r, column=c_id).value]
    why = f"zero mean: added constant c={c:g} (half smallest non-zero mean in study)"
    for cm, cs in ((c_mg, c_sg), (c_mc, c_sc)):
        set_cell(ws, r, cm, ws.cell(row=r, column=cm).value + c, why, S)
        if ws.cell(row=r, column=cs).value == 0:
            set_cell(ws, r, cs, c, why + "; zero SE set to c", S)

# Zero SE with a non-zero mean: a tiny constant would give these rows enormous
# weight, so impute the SE from the study's mean coefficient of variation.
for r in rows:
    for cm, cs in ((c_mg, c_sg), (c_mc, c_sc)):
        if ws.cell(row=r, column=cs).value == 0 and ws.cell(row=r, column=cm).value:
            study = ws.cell(row=r, column=c_id).value
            cvs = []
            for q in rows_of(study):
                for m2, s2 in ((c_mg, c_sg), (c_mc, c_sc)):
                    m_, s_ = ws.cell(row=q, column=m2).value, ws.cell(row=q, column=s2).value
                    if isinstance(m_, (int, float)) and isinstance(s_, (int, float)) and m_ > 0 and s_ > 0:
                        cvs.append(s_ / m_)
            new = ws.cell(row=r, column=cm).value * sum(cvs) / len(cvs)
            set_cell(ws, r, cs, round(new, 6),
                     f"zero SE with non-zero mean: imputed from study mean CV ({sum(cvs)/len(cvs):.3f})", S)

# --- 5. derived columns
new_cols = ["Exclusion_start_year", "Sampling_year_first", "Sampling_year_last",
            "Years_since_exclusion", "Years_since_exclusion_note",
            "Stratum", "Herbivore_origin", "Max_size_class", "Single_size_class",
            "Response_group", "Design"]
start = ws.max_column + 1
for i, name in enumerate(new_cols):
    cell = ws.cell(row=1, column=start + i, value=name)
    cell.fill = DERIVED
    cell.font = Font(bold=True)
NC = {n: start + i for i, n in enumerate(new_cols)}

STRATUM = {
    "above-ground": "Above-ground",
    "soil/litter dwellers": "Ground & litter-dwelling",
    "ground dwelling": "Ground & litter-dwelling",
    "soil/ground dwelling": "Ground & litter-dwelling",
    "above-ground and soil/litter dwellers": "Ground & litter-dwelling",
    "below-ground": "Below-ground",
}
SIZE_ORDER = ["Small", "Medium", "Large", "Mega"]


def herbivore_origin(v):
    s = (v or "").lower().replace(" ", "")
    if s in ("native(wild)", "native", "native/wild", "wild", "wild/native"):
        return "Native"
    if s.startswith("domestic"):          # includes 'Domestic - native' (native breed of livestock)
        return "Domestic"
    if s == "invasive":
        return "Invasive"
    return "Mixed"                         # Native + Domestic, All, Native/Domestic/invasive


def max_size(v):
    present = [k for k in SIZE_ORDER if k.lower() in (v or "").lower()]
    return present[-1] if present else None


def single_size(v):
    return v.strip() if (v or "").strip() in SIZE_ORDER else None


def response_group(v):
    s = (v or "").lower()
    if s in ("abundance", "density"):
        return "Abundance"
    if s == "richness":
        return "Richness"
    if "diversity" in s:
        return "Diversity"
    return None   # blank, or 'Functional richness' (single row)


def design(v):
    s = (v or "").lower()
    if "natural" in s:
        return "Natural grazed vs ungrazed"
    return "Exclosure"


for r in rows:
    dur = ws.cell(row=r, column=c_dur).value
    dury = ws.cell(row=r, column=c_dury).value
    ex = to_year(dury)
    if ex is None:
        ex = to_year(dur)
    f, l = sampling_years(ws.cell(row=r, column=c_samp).value)
    dyrs, note = duration_years(dur)
    if dyrs is None:
        dyrs, note = duration_years(dury)
    if dyrs is not None:
        yse = dyrs
    elif ex is not None and l is not None:
        yse, note = l - ex, "last sampling year - exclusion year"
    else:
        yse, note = None, "no exclusion date"
    if yse is not None and yse < 0:
        yse, note = None, f"negative ({l}-{ex}) - check"

    vals = {
        "Exclusion_start_year": ex,
        "Sampling_year_first": f,
        "Sampling_year_last": l,
        "Years_since_exclusion": yse,
        "Years_since_exclusion_note": note,
        "Stratum": STRATUM.get(str(ws.cell(row=r, column=c_strat).value or "").strip().lower()),
        "Herbivore_origin": herbivore_origin(ws.cell(row=r, column=c_herb).value),
        "Max_size_class": max_size(ws.cell(row=r, column=c_size).value),
        "Single_size_class": single_size(ws.cell(row=r, column=c_size).value),
        "Response_group": response_group(str(ws.cell(row=r, column=c_meas).value or "").strip()),
        "Design": design(ws.cell(row=r, column=H["Grazing_Type"]).value),
    }
    for k, v in vals.items():
        ws.cell(row=r, column=NC[k], value=v).fill = DERIVED

# ================================================================ ECOSYSTEM FUNCTIONS
S2 = "Ecosystem_functions"
we = wb[S2]
HE = header_map(we)
for r in range(2, we.max_row + 1):
    v = we.cell(row=r, column=HE["Study_ID"]).value
    if isinstance(v, str) and v != v.replace("\xa0", " ").strip():
        set_cell(we, r, HE["Study_ID"], v.replace("\xa0", " ").strip(), "strip whitespace/newline", S2)
ecol = we.max_column + 1
we.cell(row=1, column=ecol, value="Function_group").fill = DERIVED
we.cell(row=1, column=ecol).font = Font(bold=True)
for r in range(2, we.max_row + 1):
    m = str(we.cell(row=r, column=HE["Measured variable_cat1"]).value or "").lower()
    if we.cell(row=r, column=HE["Study_ID"]).value is None:
        continue
    if m == "pooled sd":
        g = None   # summary row ('Pooled mean') - excluded to avoid double counting
    elif "below ground" in m:
        g = "Plant biomass - below-ground"
    elif "biomass" in m:
        g = "Plant biomass - above-ground"
    elif "mineral" in m:
        g = "N mineralisation"
    elif "respiration" in m:
        g = "Respiration"
    else:
        g = None
    we.cell(row=r, column=ecol, value=g).fill = DERIVED

# ================================================================ LOG
wl = wb.create_sheet("Cleaning_log", 0)
wl.append(["Sheet", "Excel_row", "Column", "Original", "New", "Reason"])
for c in wl[1]:
    c.font = Font(bold=True)
for entry in log:
    wl.append(entry)
wl.append([])
wl.append(["Legend: yellow cells = edited in v2; blue columns = derived in v2. Original file unchanged."])
wl.append(["Flag: Study_ID #1022 contains two papers (Vaisanen et al., 2026 row 80; Andriuzzi and Wall, 2018 rows 249-253) - left as one study, please check."])
wl.append(["Flag: #1022 rows 249-253 have no 'Measured' value, so they are excluded from the response-specific models."])
wl.append(["Flag: #3610 (row 490) control SE was a broken formula (=AG491-AG490) - blanked; please re-extract."])
wl.append(["Flag: #2127 time since exclusion uses '>= 10 years of grazing management' as a lower bound - check this refers to the ungrazed treatment."])
wl.append(["Flag: n missing for #2260 (rows 162-167) and #2115 (row 491); these rows drop from the models until back-filled."])
wl.append(["Flag: Order still holds higher ranks where the order is unknown (Acari, Collembola, Gastropoda, Nematoda); Isoptera kept although termites are now placed within Blattodea."])
wl.append(["Flag: #1490 Order 'Lepidoptera; Hymenoptera (Symphyta)' covers two orders - left unchanged. #1505 bristletails have no Order (Archaeognatha or Zygentoma?) - please check."])
wl.append(["Flag: #1097 has 'Total' (whole community) rows as well as rows for each trophic group - using both would double count."])
wl.append(["Flag: possible entry errors (not changed): #3539 Hemiptera with common name 'Grasshopper'; #2115 grasshopper (Acrididae) with common name 'Leaf beetle'; spiders with Group 'Insecta' (#1125, #2260, #2394); Family typos 'Carabidiae' and 'Stapylindae' (#1190)."])
wl.column_dimensions["C"].width = 32
wl.column_dimensions["D"].width = 30
wl.column_dimensions["E"].width = 30
wl.column_dimensions["F"].width = 80

wb.save(DST)
print(f"wrote {DST}: {len(log)} edits logged")
